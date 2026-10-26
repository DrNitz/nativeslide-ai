"""
NativeSlide AI — Core DrawingML PPTX Builder
Deterministically renders structured JSON slide specs into native
100% editable OpenXML PowerPoint presentations.
"""
from __future__ import annotations

from pathlib import Path
from uuid import uuid4

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

from app.core.config import get_settings
from app.engine.themes import ThemeColors
from app.engine import vector_kit as vk
from app.models.schemas import DeckSpec, SlideLayout, SlideSpec

settings = get_settings()

# Slide canvas: 13.333 × 7.5 inches (standard widescreen 16:9)
SLIDE_W = 13.333
SLIDE_H = 7.5


class DeckBuilder:
    """
    Converts a fully-specified DeckSpec into a native python-pptx Presentation
    using only DrawingML vector primitives — no raster images embedded.
    """

    def __init__(self, spec: DeckSpec, theme: ThemeColors) -> None:
        self.spec = spec
        self.theme = theme
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)

    # ── Colour Helpers ─────────────────────────────────────────────────────

    @staticmethod
    def _hex_to_rgb(hex_str: str) -> RGBColor:
        """Convert a 6-char hex string (with or without leading #) to RGBColor."""
        h = hex_str.lstrip("#")
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
        return RGBColor(r, g, b)

    # ── Base Slide Template ────────────────────────────────────────────────

    def _base_slide(self, slide_spec: SlideSpec):
        """
        Creates a blank slide and stamps the standard chrome:
        - Full-width header bar with theme header colour
        - Thin accent stripe at the bottom of the header
        - Kicker label (small text above title)
        - Title textbox
        - Footer / slide-number textbox
        - Fade slide transition
        """
        blank_layout = self.prs.slide_layouts[6]  # blank
        slide = self.prs.slides.add_slide(blank_layout)

        t = self.theme

        # Background
        bg_fill = slide.background.fill
        bg_fill.solid()
        bg_fill.fore_color.rgb = self._hex_to_rgb(t.background_primary)

        # Header bar (full width, 1.10 inches tall)
        hdr = slide.shapes.add_shape(
            vk.MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0), Inches(SLIDE_W), Inches(1.10),
        )
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = self._hex_to_rgb(t.header_bg)
        hdr.line.fill.background()

        # Accent stripe — thin bar at the bottom edge of the header
        stripe = slide.shapes.add_shape(
            vk.MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(1.10), Inches(SLIDE_W), Inches(0.035),
        )
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = self._hex_to_rgb(t.accent_primary)
        stripe.line.fill.background()

        # Kicker textbox
        if slide_spec.kicker:
            kicker_box = slide.shapes.add_textbox(
                Inches(0.35), Inches(0.08), Inches(9.0), Inches(0.32),
            )
            vk.text(
                kicker_box,
                [{
                    "t": slide_spec.kicker.upper(),
                    "sz": 7.5,
                    "bold": True,
                    "color": self._hex_to_rgb(t.accent_primary),
                    "font": t.font_primary,
                    "space_after": 0,
                }],
                anchor=MSO_ANCHOR.MIDDLE,
                margins=(0.04, 0.0, 0.04, 0.0),
            )

        # Title textbox
        title_box = slide.shapes.add_textbox(
            Inches(0.35), Inches(0.30), Inches(12.0), Inches(0.72),
        )
        vk.text(
            title_box,
            [{
                "t": slide_spec.title,
                "sz": 20,
                "bold": True,
                "color": self._hex_to_rgb(t.white),
                "font": t.font_primary,
            }],
            anchor=MSO_ANCHOR.MIDDLE,
            margins=(0.04, 0.0, 0.04, 0.0),
        )

        # Footer / slide number
        footer_box = slide.shapes.add_textbox(
            Inches(0.35), Inches(7.25), Inches(12.5), Inches(0.22),
        )
        vk.text(
            footer_box,
            [{
                "t": f"{self.spec.title}  •  {slide_spec.slide_number}/{self.spec.total_slides}",
                "sz": 8,
                "bold": False,
                "italic": True,
                "color": self._hex_to_rgb(t.text_light),
                "font": t.font_primary,
                "align": PP_ALIGN.RIGHT,
            }],
            anchor=MSO_ANCHOR.MIDDLE,
            margins=(0.04, 0.0, 0.04, 0.0),
        )

        # Slide transition
        vk.transition(slide, kind="fade", speed="med")

        return slide

    # ── Layout Renderers ───────────────────────────────────────────────────

    def _render_metrics_callout(self, slide, spec: SlideSpec) -> None:
        """Render up to 3 metric callout cards in a top row."""
        if not spec.metrics:
            return

        t = self.theme
        metrics = spec.metrics[:3]
        n = len(metrics)
        card_w = (SLIDE_W - 0.70 - (n - 1) * 0.22) / n
        card_h = 4.90
        y_start = 1.30

        for i, m in enumerate(metrics):
            x = 0.35 + i * (card_w + 0.22)

            # Card background
            c = vk.card(
                slide, x, y_start, card_w, card_h,
                fill=self._hex_to_rgb(t.card_fill),
                line=self._hex_to_rgb(t.card_border),
                shadow=(8, 4, 18),
            )
            c.text_frame.text = ""

            accent_hex = m.accent_color or t.accent_primary
            accent_rgb = self._hex_to_rgb(accent_hex)

            # Top accent bar on card
            accent_bar = slide.shapes.add_shape(
                vk.MSO_SHAPE.RECTANGLE,
                Inches(x), Inches(y_start), Inches(card_w), Inches(0.06),
            )
            accent_bar.fill.solid()
            accent_bar.fill.fore_color.rgb = accent_rgb
            accent_bar.line.fill.background()

            # Large value
            val_box = slide.shapes.add_textbox(
                Inches(x + 0.18), Inches(y_start + 0.20),
                Inches(card_w - 0.36), Inches(1.10),
            )
            vk.text(
                val_box,
                [{"t": m.value, "sz": 26, "bold": True,
                  "color": accent_rgb, "font": t.font_primary}],
                anchor=MSO_ANCHOR.MIDDLE,
                margins=(0.05, 0.02, 0.05, 0.02),
            )

            # Label
            label_box = slide.shapes.add_textbox(
                Inches(x + 0.18), Inches(y_start + 1.35),
                Inches(card_w - 0.36), Inches(0.38),
            )
            vk.text(
                label_box,
                [{"t": m.label, "sz": 9.5, "bold": True,
                  "color": self._hex_to_rgb(t.text_dark), "font": t.font_primary}],
                anchor=MSO_ANCHOR.TOP,
                margins=(0.05, 0.02, 0.05, 0.02),
            )

            # Description
            desc_box = slide.shapes.add_textbox(
                Inches(x + 0.18), Inches(y_start + 1.80),
                Inches(card_w - 0.36), Inches(2.80),
            )
            vk.text(
                desc_box,
                [{"t": m.description, "sz": 10, "bold": False,
                  "color": self._hex_to_rgb(t.text_medium), "font": t.font_primary,
                  "line": 1.3}],
                anchor=MSO_ANCHOR.TOP,
                margins=(0.05, 0.02, 0.05, 0.02),
            )

    def _render_table_matrix(self, slide, spec: SlideSpec) -> None:
        """Render a native <a:tbl> table from spec.table data."""
        if not spec.table:
            return

        t = self.theme
        td = spec.table
        n_cols = len(td.headers)
        n_rows = len(td.rows)

        tbl_x = Inches(0.35)
        tbl_y = Inches(1.30)
        tbl_w = Inches(SLIDE_W - 0.70)
        tbl_h = Inches(SLIDE_H - 1.65)

        table = slide.shapes.add_table(
            n_rows + 1, n_cols, tbl_x, tbl_y, tbl_w, tbl_h,
        ).table

        # Column widths
        col_total = tbl_w
        if td.column_widths and len(td.column_widths) == n_cols:
            for ci, frac in enumerate(td.column_widths):
                table.columns[ci].width = Emu(int(col_total * frac))
        else:
            equal_w = col_total // n_cols
            for ci in range(n_cols):
                table.columns[ci].width = equal_w

        header_bg = self._hex_to_rgb(t.table_header_bg)
        header_fg = self._hex_to_rgb(t.white)
        alt_bg = self._hex_to_rgb(t.table_row_alt)
        body_fg = self._hex_to_rgb(t.text_dark)

        # Header row
        for ci, hdr_text in enumerate(td.headers):
            cell = table.cell(0, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_bg
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = hdr_text
            r.font.name = t.font_primary
            r.font.size = Pt(9)
            r.font.bold = True
            r.font.color.rgb = header_fg

        # Data rows
        for ri, row_data in enumerate(td.rows):
            is_alt = ri % 2 == 1
            for ci, cell_text in enumerate(row_data):
                if ci >= n_cols:
                    break
                cell = table.cell(ri + 1, ci)
                if is_alt:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = alt_bg
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = self._hex_to_rgb(t.card_fill)
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT
                r = p.add_run()
                r.text = cell_text
                r.font.name = t.font_primary
                r.font.size = Pt(8.5)
                r.font.color.rgb = body_fg

    def _render_cards_3x2(self, slide, spec: SlideSpec) -> None:
        """Render up to 6 cards in a 3-column × 2-row grid."""
        if not spec.cards:
            return

        t = self.theme
        cards = spec.cards[:6]
        cols = 3
        rows = 2
        card_w = (SLIDE_W - 0.70 - (cols - 1) * 0.20) / cols
        card_h = (SLIDE_H - 1.50 - (rows - 1) * 0.18) / rows
        x0, y0 = 0.35, 1.32

        for idx, card_item in enumerate(cards):
            col = idx % cols
            row = idx // cols
            x = x0 + col * (card_w + 0.20)
            y = y0 + row * (card_h + 0.18)

            accent_hex = card_item.accent_color or t.accent_primary
            accent_rgb = self._hex_to_rgb(accent_hex)

            # Card body
            vk.card(
                slide, x, y, card_w, card_h,
                fill=self._hex_to_rgb(t.card_fill),
                line=self._hex_to_rgb(t.card_border),
                shadow=(6, 3, 14),
            )

            # Top accent bar
            ab = slide.shapes.add_shape(
                vk.MSO_SHAPE.RECTANGLE,
                Inches(x), Inches(y), Inches(card_w), Inches(0.05),
            )
            ab.fill.solid()
            ab.fill.fore_color.rgb = accent_rgb
            ab.line.fill.background()

            # Title pill badge
            vk.pill(
                slide, x + 0.10, y + 0.10, card_w - 0.20, 0.26,
                fill=accent_rgb,
                label=card_item.title,
                sz=7.5,
            )

            # Body text
            body_box = slide.shapes.add_textbox(
                Inches(x + 0.10), Inches(y + 0.42),
                Inches(card_w - 0.20), Inches(card_h - 0.55),
            )
            vk.text(
                body_box,
                [{"t": card_item.body, "sz": 9, "bold": False,
                  "color": self._hex_to_rgb(t.text_medium), "font": t.font_primary,
                  "line": 1.25}],
                anchor=MSO_ANCHOR.TOP,
                margins=(0.04, 0.02, 0.04, 0.02),
            )

    def _render_comparison_2col(self, slide, spec: SlideSpec) -> None:
        """Render 2 wide comparison cards side by side."""
        if not spec.cards:
            return

        t = self.theme
        cards = spec.cards[:2]
        card_w = 5.85
        card_h = SLIDE_H - 1.50
        y = 1.30

        positions = [0.35, 0.35 + card_w + 0.28]

        for i, (card_item, x) in enumerate(zip(cards, positions)):
            accent_hex = card_item.accent_color or (
                t.accent_primary if i == 0 else t.accent_secondary
            )
            accent_rgb = self._hex_to_rgb(accent_hex)

            vk.card(
                slide, x, y, card_w, card_h,
                fill=self._hex_to_rgb(t.card_fill),
                line=self._hex_to_rgb(t.card_border),
                shadow=(8, 4, 18),
            )

            # Accent bar
            ab = slide.shapes.add_shape(
                vk.MSO_SHAPE.RECTANGLE,
                Inches(x), Inches(y), Inches(card_w), Inches(0.06),
            )
            ab.fill.solid()
            ab.fill.fore_color.rgb = accent_rgb
            ab.line.fill.background()

            # Title pill
            vk.pill(
                slide, x + 0.15, y + 0.12, card_w - 0.30, 0.30,
                fill=accent_rgb,
                label=card_item.title,
                sz=8.5,
            )

            # Subtitle
            if card_item.subtitle:
                sub_box = slide.shapes.add_textbox(
                    Inches(x + 0.15), Inches(y + 0.50),
                    Inches(card_w - 0.30), Inches(0.30),
                )
                vk.text(
                    sub_box,
                    [{"t": card_item.subtitle, "sz": 9, "bold": True, "italic": True,
                      "color": accent_rgb, "font": t.font_primary}],
                    anchor=MSO_ANCHOR.MIDDLE,
                    margins=(0.04, 0.0, 0.04, 0.0),
                )

            # Body
            body_y = y + 0.85 if card_item.subtitle else y + 0.52
            body_box = slide.shapes.add_textbox(
                Inches(x + 0.15), Inches(body_y),
                Inches(card_w - 0.30), Inches(card_h - body_y + y - 0.20),
            )
            vk.text(
                body_box,
                [{"t": card_item.body, "sz": 10, "bold": False,
                  "color": self._hex_to_rgb(t.text_medium), "font": t.font_primary,
                  "line": 1.35}],
                anchor=MSO_ANCHOR.TOP,
                margins=(0.05, 0.03, 0.05, 0.03),
            )

    def _render_title_hero(self, slide, spec: SlideSpec) -> None:
        """Render a full-slide hero / opening slide with body text."""
        t = self.theme
        if not spec.body_text:
            return

        body_box = slide.shapes.add_textbox(
            Inches(0.70), Inches(1.50), Inches(SLIDE_W - 1.40), Inches(5.50),
        )
        vk.text(
            body_box,
            [{"t": spec.body_text, "sz": 13, "bold": False,
              "color": self._hex_to_rgb(t.text_dark), "font": t.font_primary,
              "line": 1.55}],
            anchor=MSO_ANCHOR.MIDDLE,
            margins=(0.10, 0.05, 0.10, 0.05),
        )

    def _render_section_divider(self, slide, spec: SlideSpec) -> None:
        """Render a full-width section divider slide."""
        t = self.theme

        # Large accent bar across the middle
        div = slide.shapes.add_shape(
            vk.MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(3.40), Inches(SLIDE_W), Inches(0.10),
        )
        div.fill.solid()
        div.fill.fore_color.rgb = self._hex_to_rgb(t.accent_primary)
        div.line.fill.background()

        if spec.body_text:
            body_box = slide.shapes.add_textbox(
                Inches(1.50), Inches(3.60), Inches(SLIDE_W - 3.0), Inches(2.50),
            )
            vk.text(
                body_box,
                [{"t": spec.body_text, "sz": 11, "bold": False, "italic": True,
                  "color": self._hex_to_rgb(t.text_medium), "font": t.font_primary,
                  "align": PP_ALIGN.CENTER, "line": 1.45}],
                anchor=MSO_ANCHOR.TOP,
                margins=(0.08, 0.04, 0.08, 0.04),
            )

    def _render_process_steps(self, slide, spec: SlideSpec) -> None:
        """Render process steps as a horizontal flow of numbered pill cards."""
        if not spec.cards:
            return

        t = self.theme
        steps = spec.cards[:5]
        n = len(steps)
        card_w = (SLIDE_W - 0.70 - (n - 1) * 0.18) / n
        card_h = 4.80
        y = 1.35

        for i, step in enumerate(steps):
            x = 0.35 + i * (card_w + 0.18)
            accent_rgb = self._hex_to_rgb(step.accent_color or t.accent_primary)

            vk.card(
                slide, x, y, card_w, card_h,
                fill=self._hex_to_rgb(t.card_fill),
                line=self._hex_to_rgb(t.card_border),
                shadow=(5, 3, 12),
            )

            # Step number pill
            vk.pill(
                slide, x + card_w / 2 - 0.22, y + 0.12, 0.44, 0.44,
                fill=accent_rgb,
                label=str(i + 1),
                sz=11,
            )

            # Step title
            title_box = slide.shapes.add_textbox(
                Inches(x + 0.08), Inches(y + 0.64),
                Inches(card_w - 0.16), Inches(0.45),
            )
            vk.text(
                title_box,
                [{"t": step.title, "sz": 9, "bold": True,
                  "color": self._hex_to_rgb(t.text_dark), "font": t.font_primary,
                  "align": PP_ALIGN.CENTER}],
                anchor=MSO_ANCHOR.MIDDLE,
                margins=(0.04, 0.0, 0.04, 0.0),
            )

            # Body
            body_box = slide.shapes.add_textbox(
                Inches(x + 0.08), Inches(y + 1.14),
                Inches(card_w - 0.16), Inches(card_h - 1.28),
            )
            vk.text(
                body_box,
                [{"t": step.body, "sz": 8.5, "bold": False,
                  "color": self._hex_to_rgb(t.text_medium), "font": t.font_primary,
                  "line": 1.25}],
                anchor=MSO_ANCHOR.TOP,
                margins=(0.04, 0.02, 0.04, 0.02),
            )

            # Arrow connector between steps
            if i < n - 1:
                arrow_x = x + card_w + 0.01
                vk.connector(
                    slide,
                    arrow_x, y + card_h / 2,
                    arrow_x + 0.16, y + card_h / 2,
                    color=self._hex_to_rgb(t.accent_primary),
                    width=1.5,
                )

    # ── Speaker Notes ──────────────────────────────────────────────────────

    @staticmethod
    def _set_notes(slide, notes_text: str) -> None:
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    # ── Build Entry Point ──────────────────────────────────────────────────

    def build(self) -> Path:
        """
        Iterate over all slides in the DeckSpec, render each one, attach
        speaker notes, then save the Presentation to OUTPUT_DIR.

        Returns:
            Path to the saved .pptx file.
        """
        dispatch = {
            SlideLayout.METRICS_CALLOUT: self._render_metrics_callout,
            SlideLayout.TABLE_MATRIX: self._render_table_matrix,
            SlideLayout.CARDS_3X2: self._render_cards_3x2,
            SlideLayout.CASE_SCENARIOS: self._render_cards_3x2,
            SlideLayout.COMPARISON_2COL: self._render_comparison_2col,
            SlideLayout.PROCESS_STEPS: self._render_process_steps,
            SlideLayout.TITLE_HERO: self._render_title_hero,
            SlideLayout.SECTION_DIVIDER: self._render_section_divider,
        }

        for slide_spec in self.spec.slides:
            slide = self._base_slide(slide_spec)

            renderer = dispatch.get(slide_spec.layout)
            if renderer:
                renderer(slide, slide_spec)

            self._set_notes(slide, slide_spec.speaker_notes)

        output_dir = settings.OUTPUT_DIR
        output_dir.mkdir(parents=True, exist_ok=True)
        out_path = output_dir / f"{uuid4()}.pptx"
        self.prs.save(str(out_path))
        return out_path
